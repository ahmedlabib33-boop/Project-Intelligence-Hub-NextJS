from __future__ import annotations

import json
import math
import re
from datetime import datetime, timezone
from difflib import SequenceMatcher
from io import BytesIO
from typing import Any, Dict, Iterable, List, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.cpm import calculate_cpm
from core.models import Activity, Calendar, ProjectSchedule, Relationship, parse_dt
from core.xer import parse_xer_bytes

P6_WARNING = (
    "Pre-P6 planning model only. Import into the intended Primavera P6 environment, "
    "assign approved native calendars, schedule in P6, then validate dates, float, "
    "constraints, open ends and longest path before reliance."
)

FIXED_QUESTIONS = [
    {"id": "project_id", "question": "What is the unique project ID?", "critical": True},
    {"id": "commencement_date", "question": "What is the approved commencement / NTP date?", "critical": True},
    {"id": "data_date", "question": "What is the schedule data date?", "critical": True},
    {"id": "target_duration_days", "question": "What is the target duration in working days?", "critical": True},
    {"id": "hours_per_day", "question": "How many working hours are in one day?", "critical": True},
    {"id": "days_per_week", "question": "How many working days are in one week?", "critical": True},
    {"id": "holiday_dates", "question": "Which approved dates are non-working holidays?", "critical": False},
    {"id": "default_wbs", "question": "Which approved WBS receives activities without a WBS?", "critical": True},
    {"id": "logic_mode", "question": "Use evidence logic, or approve sequential source-order logic?", "critical": True},
    {"id": "conflicts_resolved", "question": "Are all Analyzer schedule conflicts resolved?", "critical": True},
]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _norm(value: Any) -> str:
    return " ".join(re.findall(r"[a-z0-9\u0600-\u06ff]+", _text(value).lower()))


def parse_payload(value: str | Mapping[str, Any] | None, label: str) -> Dict[str, Any]:
    if isinstance(value, Mapping):
        return dict(value)
    try:
        decoded = json.loads(value or "{}")
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid {label} JSON at line {exc.lineno}") from None
    if not isinstance(decoded, dict):
        raise ValueError(f"{label} must be a JSON object")
    return decoded


def _rows(payload: Mapping[str, Any]) -> List[Dict[str, Any]]:
    return [dict(row) for row in payload.get("rows") or [] if isinstance(row, Mapping)]


def _approved(row: Mapping[str, Any]) -> bool:
    status = _text(row.get("approval_status")).upper().replace(" ", "_")
    return _text(row.get("working_status") or "ACTIVE").upper() != "EXCLUDED" and status in {"APPROVED", "SELECTED", "INCLUDED"}


def _activity_name(row: Mapping[str, Any]) -> str:
    return _text(row.get("display_name") or row.get("source_activity_name") or row.get("activity_name"))


def _best_rate(activity: Mapping[str, Any], rates: Sequence[Mapping[str, Any]]) -> Mapping[str, Any] | None:
    name = _norm(_activity_name(activity))
    scored = []
    for rate in rates:
        candidate = _norm(rate.get("display_name") or rate.get("original_activity_name") or rate.get("activity_name"))
        if not name or not candidate:
            continue
        score = 1.0 if name == candidate else SequenceMatcher(None, name, candidate).ratio()
        if name in candidate or candidate in name:
            score = max(score, 0.9)
        scored.append((score, rate))
    if not scored:
        return None
    score, rate = max(scored, key=lambda item: item[0])
    return rate if score >= 0.58 else None


def _section(analyzer: Mapping[str, Any], key: str) -> List[Dict[str, Any]]:
    section = analyzer.get("schedule_builder_output") or {}
    value = section.get(key) if isinstance(section, Mapping) else []
    return [dict(row) for row in value or [] if isinstance(row, Mapping)]


def _holiday_dates(value: Any) -> List[str]:
    candidates = value if isinstance(value, list) else re.split(r"[,;\n]+", _text(value))
    return sorted({parsed.date().isoformat() for item in candidates if (parsed := parse_dt(item))})


def build_schedule_development(
    analyzer_payload: str | Mapping[str, Any] | None,
    productivity_payload: str | Mapping[str, Any] | None,
    activity_payload: str | Mapping[str, Any] | None,
    answers_payload: str | Mapping[str, Any] | None,
    crew_overrides_payload: str | Mapping[str, Any] | None = None,
) -> Dict[str, Any]:
    analyzer = parse_payload(analyzer_payload, "Analyzer payload")
    productivity = parse_payload(productivity_payload, "Productivity payload")
    activity_library = parse_payload(activity_payload, "Activity List payload")
    answers = parse_payload(answers_payload, "Schedule answers")
    crew_overrides = parse_payload(crew_overrides_payload, "Crew overrides")
    blockers: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    questions = [{**q, "answer": answers.get(q["id"]), "status": "ANSWERED" if answers.get(q["id"]) not in (None, "", False) else "MISSING"} for q in FIXED_QUESTIONS]

    project_id = _text(answers.get("project_id"))
    project_name = _text(answers.get("project_name") or analyzer.get("project") or project_id)
    commencement, data_date = parse_dt(answers.get("commencement_date")), parse_dt(answers.get("data_date"))
    target = _number(answers.get("target_duration_days"))
    hours, days = _number(answers.get("hours_per_day")), _number(answers.get("days_per_week"))
    default_wbs, logic_mode = _text(answers.get("default_wbs")), _text(answers.get("logic_mode")).upper()
    required = {
        "project_id": bool(project_id), "commencement_date": bool(commencement), "data_date": bool(data_date),
        "target_duration_days": bool(target and target > 0), "hours_per_day": bool(hours and 0 < hours <= 24),
        "days_per_week": bool(days and 1 <= days <= 7), "default_wbs": bool(default_wbs),
        "logic_mode": logic_mode in {"EVIDENCE", "SEQUENTIAL_SOURCE_ORDER"},
    }
    for field, valid in required.items():
        if not valid:
            blockers.append({"type": "MISSING_FIXED_ANSWER", "field": field, "message": f"Answer the fixed question: {field.replace('_', ' ')}."})
    conflicts = analyzer.get("conflicts") or []
    if conflicts and not bool(answers.get("conflicts_resolved")):
        blockers.append({"type": "UNRESOLVED_ANALYZER_CONFLICTS", "field": "conflicts_resolved", "message": f"Resolve and confirm {len(conflicts)} Analyzer conflict(s)."})

    activities = sorted([row for row in _rows(activity_library) if _approved(row)], key=lambda row: (_number(row.get("source_sequence_number")) or 10**12, _activity_name(row)))
    rates = [row for row in _rows(productivity) if _approved(row)]
    if not activities:
        blockers.append({"type": "NO_APPROVED_ACTIVITIES", "field": "Activity List", "message": "Approve activities in the Activity List tab."})
    if not rates:
        blockers.append({"type": "NO_APPROVED_PRODUCTIVITY", "field": "Productivity Rates", "message": "Approve productivity rows with quantity, selected productivity and crews."})

    crew_plan, schedule_rows, seen_ids = [], [], set()
    for index, row in enumerate(activities, start=1):
        activity_id = _text(row.get("activity_id") or row.get("id") or f"ACT-{index:05d}")
        if activity_id in seen_ids:
            activity_id = f"{activity_id}-{index}"
        seen_ids.add(activity_id)
        name, rate = _activity_name(row), _best_rate(row, rates)
        quantity, production = _number((rate or {}).get("quantity")), _number((rate or {}).get("selected_productivity"))
        crews = _number(crew_overrides.get(activity_id)) or _number((rate or {}).get("number_of_crews"))
        duration = _number(row.get("duration_days"))
        if duration is None and quantity is not None and production and production > 0 and crews and crews > 0:
            duration = max(1.0, float(math.ceil(quantity / (production * crews))))
        if not name:
            blockers.append({"type": "ACTIVITY_NAME_MISSING", "activity_id": activity_id, "message": "Approved activity has no display name."})
        if duration is None:
            missing = (["quantity"] if quantity is None else []) + (["selected productivity"] if not production or production <= 0 else []) + (["number of crews"] if not crews or crews <= 0 else [])
            blockers.append({"type": "DURATION_BASIS_MISSING", "activity_id": activity_id, "message": f"{name or activity_id}: missing {', '.join(missing) or 'approved duration basis'}."})
        recommended = max(1, math.ceil(quantity / (production * target))) if quantity is not None and production and target else None
        schedule_row = {
            "activity_id": activity_id, "activity_name": name, "wbs_id": _text(row.get("wbs") or row.get("wbs_code") or default_wbs),
            "discipline": _text(row.get("discipline")), "work_package": _text(row.get("work_package")), "quantity": quantity,
            "uom": _text((rate or {}).get("normalized_uom") or (rate or {}).get("original_uom") or row.get("uom")),
            "selected_productivity": production, "number_of_crews": crews, "duration_days": duration,
            "source_activity": _text(row.get("source_activity_name")), "productivity_source": _text((rate or {}).get("source_document")),
            "productivity_location": _text((rate or {}).get("source_location")), "activity_approval": _text(row.get("approval_status")),
            "rate_approval": _text((rate or {}).get("approval_status")),
        }
        schedule_rows.append(schedule_row)
        crew_plan.append({"activity_id": activity_id, "activity_name": name, "quantity": quantity, "uom": schedule_row["uom"], "daily_productivity_per_crew": production, "recommended_crews_for_target": recommended, "selected_crews": crews, "calculated_duration_days": duration, "target_project_duration_days": target})

    relationships, activity_ids = [], {row["activity_id"] for row in schedule_rows}
    if logic_mode == "EVIDENCE":
        for index, row in enumerate(_section(analyzer, "relationships"), start=1):
            pred, succ = _text(row.get("Predecessor") or row.get("predecessor")), _text(row.get("Successor") or row.get("successor"))
            rel_type = _text(row.get("Relationship") or row.get("relationship") or "FS").upper().replace("PR_", "")
            lag = _number(row.get("Lag") or row.get("lag")) or 0.0
            if pred in activity_ids and succ in activity_ids:
                relationships.append({"id": f"REL-{index:05d}", "predecessor_id": pred, "successor_id": succ, "rel_type": rel_type if rel_type in {"FS", "SS", "FF", "SF"} else "FS", "lag_days": lag, "basis": "Analyzer evidence", "source": _text(row.get("Source"))})
        if len(schedule_rows) > 1 and not relationships:
            blockers.append({"type": "NO_MATCHING_EVIDENCE_LOGIC", "field": "logic_mode", "message": "Analyzer logic does not match approved Activity List IDs. Supply matching logic or approve sequential source order."})
    elif logic_mode == "SEQUENTIAL_SOURCE_ORDER":
        relationships = [{"id": f"REL-{index:05d}", "predecessor_id": schedule_rows[index - 1]["activity_id"], "successor_id": schedule_rows[index]["activity_id"], "rel_type": "FS", "lag_days": 0.0, "basis": "User-approved sequential source-order assumption", "source": "Schedule Development fixed answer"} for index in range(1, len(schedule_rows))]
        warnings.append({"type": "APPROVED_ASSUMPTION", "message": "Finish-to-Start source-order logic must be replaced by approved project logic before baseline acceptance."})

    steps = [
        {"step": 1, "name": "Analyzer evidence gate", "status": "COMPLETE" if analyzer else "BLOCKED", "detail": "Analyzer result received" if analyzer else "Run Analyzer first"},
        {"step": 2, "name": "Activity scope selection", "status": "COMPLETE" if activities else "BLOCKED", "detail": f"{len(activities)} approved activities"},
        {"step": 3, "name": "Productivity and crew matching", "status": "COMPLETE" if rates else "BLOCKED", "detail": f"{len(rates)} approved rate rows"},
        {"step": 4, "name": "Duration calculation", "status": "COMPLETE" if schedule_rows and all(row["duration_days"] is not None for row in schedule_rows) else "BLOCKED", "detail": "Quantity / (daily productivity per crew × crews), rounded up"},
        {"step": 5, "name": "Calendar setup", "status": "COMPLETE" if required["hours_per_day"] and required["days_per_week"] else "BLOCKED", "detail": f"{hours or 'N/A'} h/day; {days or 'N/A'} days/week"},
        {"step": 6, "name": "Logic network", "status": "COMPLETE" if len(schedule_rows) <= 1 or relationships else "BLOCKED", "detail": f"{len(relationships)} relationships"},
        {"step": 7, "name": "Shadow CPM calculation", "status": "WAITING", "detail": "Runs after critical blockers close"},
        {"step": 8, "name": "Excel and XER staging", "status": "WAITING", "detail": "Native Primavera P6 import and recalculation remain mandatory"},
    ]
    output: Dict[str, Any] = {
        "engine": "Generic Schedule Development", "status": "BLOCKED" if blockers else "READY_TO_CALCULATE", "fixed_questions": questions,
        "blockers": blockers, "warnings": warnings, "construction_steps": steps, "crew_plan": crew_plan, "activities": schedule_rows,
        "relationships": relationships, "wbs": _section(analyzer, "wbs") or ([{"WBS": default_wbs, "Source": "Fixed answer", "Status": "APPROVED INPUT"}] if default_wbs else []),
        "source_counts": {"analyzer_documents": len(analyzer.get("document_register") or []), "approved_activities": len(activities), "approved_productivity_rows": len(rates), "relationships": len(relationships)},
        "native_p6_verification_required": True, "warning": P6_WARNING,
    }
    if blockers:
        return output

    schedule = ProjectSchedule(
        project_id=project_id, project_name=project_name or project_id, data_date=data_date,
        calendars={"CAL-DEFAULT": Calendar(id="CAL-DEFAULT", name="Approved Schedule Development Calendar", working_weekdays=list(range(int(days or 5))), hours_per_day=float(hours or 8), non_working_dates=_holiday_dates(answers.get("holiday_dates")))},
        wbs={default_wbs: default_wbs}, metadata={"default_calendar_id": "CAL-DEFAULT", "default_hours_per_day": hours, "target_duration_days": target, "required_finish_date": _text(answers.get("required_finish_date")), "source": "Analyzer + approved Productivity Rates + approved Activity List + fixed answers", "native_p6_verification_required": True},
    )
    for row in schedule_rows:
        schedule.activities[row["activity_id"]] = Activity(id=row["activity_id"], name=row["activity_name"], wbs_id=row["wbs_id"], calendar_id="CAL-DEFAULT", original_duration=float(row["duration_days"]), remaining_duration=float(row["duration_days"]), planned_start=commencement, source={key: value for key, value in row.items() if key not in {"activity_id", "activity_name"}})
    for row in relationships:
        schedule.relationships.append(Relationship(predecessor_id=row["predecessor_id"], successor_id=row["successor_id"], rel_type=row["rel_type"], lag=float(row["lag_days"]), id=row["id"], source={"basis": row["basis"], "source": row["source"]}))
    try:
        calculate_cpm(schedule)
    except Exception as exc:
        output["status"] = "BLOCKED"
        output["blockers"].append({"type": "SHADOW_CPM_FAILED", "message": f"{type(exc).__name__}: {exc}"})
        return output
    for row in schedule_rows:
        current = schedule.activities[row["activity_id"]]
        row.update({"forecast_start": current.forecast_start.isoformat() if current.forecast_start else "", "forecast_finish": current.forecast_finish.isoformat() if current.forecast_finish else "", "total_float_days": current.total_float, "critical": current.critical, "longest_path": current.longest_path})
    output["status"] = "PRE_P6_SCHEDULE_BUILT"
    output["summary"] = {"project_id": project_id, "project_name": project_name or project_id, "activity_count": len(schedule.activities), "relationship_count": len(schedule.relationships), "commencement_date": commencement.date().isoformat(), "data_date": data_date.date().isoformat(), "target_duration_days": target, "forecast_finish": max((item.forecast_finish for item in schedule.activities.values() if item.forecast_finish), default=None), "critical_count": sum(1 for item in schedule.activities.values() if item.critical)}
    output["normalized_schedule"] = schedule.to_dict(include_raw_tables=False)
    output["construction_steps"][-2].update(status="COMPLETE", detail=f"Shadow CPM calculated for {len(schedule.activities)} activities; native P6 remains authoritative")
    output["construction_steps"][-1]["status"] = "READY"
    return output


def _write_sheet(workbook: Workbook, title: str, source_rows: Iterable[Mapping[str, Any]]) -> None:
    records, columns = [dict(row) for row in source_rows], []
    for row in records:
        columns.extend(key for key in row if key not in columns)
    if not columns:
        columns, records = ["status"], [{"status": "No records"}]
    sheet = workbook.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", title)[:31])
    for index, column in enumerate(columns, 1):
        cell = sheet.cell(1, index, column.replace("_", " ").title())
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="073B5C"), Alignment(wrap_text=True)
    for row_index, row in enumerate(records, 2):
        for column_index, column in enumerate(columns, 1):
            value = row.get(column)
            sheet.cell(row_index, column_index, json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value)
    sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in list(cells)[:100]) + 2, 12), 48)


def build_schedule_development_xlsx(result: Mapping[str, Any]) -> bytes:
    if result.get("status") != "PRE_P6_SCHEDULE_BUILT":
        raise ValueError("Close all Schedule Development blockers before export")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, source_rows in (("Executive Summary", [dict(result.get("summary") or {}), {"warning": P6_WARNING}]), ("Fixed Questions", result.get("fixed_questions") or []), ("Construction Steps", result.get("construction_steps") or []), ("WBS", result.get("wbs") or []), ("Activities", result.get("activities") or []), ("Logic", result.get("relationships") or []), ("Crew Plan", result.get("crew_plan") or []), ("Warnings", result.get("warnings") or [])):
        _write_sheet(workbook, title, source_rows)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _xer_cell(value: Any) -> str:
    return _text(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _xer_table(name: str, fields: Sequence[str], source_rows: Sequence[Mapping[str, Any]]) -> List[str]:
    return [f"%T\t{name}", "%F\t" + "\t".join(fields)] + ["%R\t" + "\t".join(_xer_cell(row.get(field)) for field in fields) for row in source_rows]


def build_schedule_development_xer(result: Mapping[str, Any]) -> bytes:
    if result.get("status") != "PRE_P6_SCHEDULE_BUILT":
        raise ValueError("Close all Schedule Development blockers before export")
    summary = dict(result.get("summary") or {})
    activities, relationships = [dict(row) for row in result.get("activities") or []], [dict(row) for row in result.get("relationships") or []]
    project_id, project_internal, calendar_internal = _xer_cell(summary.get("project_id") or "PROJECT"), "1", "1"
    questions = result.get("fixed_questions") or []
    hours = _number(next((item.get("answer") for item in questions if item.get("id") == "hours_per_day"), 8)) or 8
    days = _number(next((item.get("answer") for item in questions if item.get("id") == "days_per_week"), 5)) or 5
    lines = [f"ERMHDR\t16.2\t{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}\tProject Intelligence Hub\tSchedule Development\tAdmin\tUSD"]
    lines += _xer_table("PROJECT", ["proj_id", "proj_short_name", "proj_name", "last_recalc_date", "plan_start_date", "plan_end_date", "clndr_id"], [{"proj_id": project_internal, "proj_short_name": project_id, "proj_name": summary.get("project_name"), "last_recalc_date": summary.get("data_date"), "plan_start_date": summary.get("commencement_date"), "plan_end_date": summary.get("forecast_finish"), "clndr_id": calendar_internal}])
    lines += _xer_table("CALENDAR", ["clndr_id", "clndr_name", "day_hr_cnt", "week_hr_cnt"], [{"clndr_id": calendar_internal, "clndr_name": "Approved Schedule Development Calendar", "day_hr_cnt": hours, "week_hr_cnt": hours * days}])
    wbs_codes = list(dict.fromkeys(_text(row.get("wbs_id")) or "WBS" for row in activities))
    wbs_ids = {code: str(index + 1) for index, code in enumerate(wbs_codes)}
    lines += _xer_table("PROJWBS", ["wbs_id", "proj_id", "wbs_short_name", "wbs_name", "parent_wbs_id", "seq_num"], [{"wbs_id": wbs_ids[code], "proj_id": project_internal, "wbs_short_name": code, "wbs_name": code, "parent_wbs_id": "", "seq_num": index + 1} for index, code in enumerate(wbs_codes)])
    task_ids = {row["activity_id"]: str(index + 1) for index, row in enumerate(activities)}
    task_rows = [{"task_id": task_ids[row["activity_id"]], "proj_id": project_internal, "wbs_id": wbs_ids[_text(row.get("wbs_id")) or "WBS"], "clndr_id": calendar_internal, "task_code": row["activity_id"], "task_name": row["activity_name"], "task_type": "TT_Task", "status_code": "TK_NotStart", "target_drtn_hr_cnt": float(row["duration_days"]) * hours, "remain_drtn_hr_cnt": float(row["duration_days"]) * hours, "target_start_date": _text(row.get("forecast_start")), "target_end_date": _text(row.get("forecast_finish"))} for row in activities]
    lines += _xer_table("TASK", ["task_id", "proj_id", "wbs_id", "clndr_id", "task_code", "task_name", "task_type", "status_code", "target_drtn_hr_cnt", "remain_drtn_hr_cnt", "target_start_date", "target_end_date"], task_rows)
    pred_rows = [{"task_pred_id": index + 1, "task_id": task_ids[row["successor_id"]], "pred_task_id": task_ids[row["predecessor_id"]], "pred_type": f"PR_{row['rel_type']}", "lag_hr_cnt": float(row["lag_days"]) * hours} for index, row in enumerate(relationships) if row["predecessor_id"] in task_ids and row["successor_id"] in task_ids]
    lines += _xer_table("TASKPRED", ["task_pred_id", "task_id", "pred_task_id", "pred_type", "lag_hr_cnt"], pred_rows)
    payload = ("\r\n".join(lines + ["%E"]) + "\r\n").encode("utf-8")
    reparsed = parse_xer_bytes(payload, default_hours_per_day=hours, project_id=project_internal)
    if len(reparsed.activities) != len(activities) or len(reparsed.relationships) != len(pred_rows):
        raise ValueError("Generated XER failed internal round-trip count validation")
    return payload


def project_schedule_to_development_result(schedule: ProjectSchedule) -> Dict[str, Any]:
    """Adapt a governed working/revised model to the neutral R16.2 staging exporter.

    This deliberately carries schedule facts only. It does not represent native
    Primavera recalculation or approval, and the exported file retains the same
    mandatory P6 warning as the initial Schedule Development export.
    """
    calculate_cpm(schedule)
    calendar = schedule.get_calendar()
    activities = [
        {
            "activity_id": activity.id,
            "activity_name": activity.name,
            "wbs_id": activity.wbs_id or "WBS",
            "duration_days": float(activity.remaining_duration),
            "forecast_start": activity.forecast_start,
            "forecast_finish": activity.forecast_finish,
        }
        for activity in schedule.activities.values()
    ]
    relationships = [
        {
            "predecessor_id": relationship.predecessor_id,
            "successor_id": relationship.successor_id,
            "rel_type": relationship.rel_type,
            "lag_days": float(relationship.lag),
        }
        for relationship in schedule.relationships
    ]
    forecast_dates = [row["forecast_finish"] for row in activities if row["forecast_finish"]]
    commencement_dates = [row["forecast_start"] for row in activities if row["forecast_start"]]
    return {
        "status": "PRE_P6_SCHEDULE_BUILT",
        "summary": {
            "project_id": schedule.project_id,
            "project_name": schedule.project_name,
            "data_date": schedule.data_date,
            "commencement_date": min(commencement_dates) if commencement_dates else schedule.data_date,
            "forecast_finish": max(forecast_dates) if forecast_dates else None,
        },
        "fixed_questions": [
            {"id": "hours_per_day", "answer": calendar.hours_per_day},
            {"id": "days_per_week", "answer": len(calendar.working_weekdays)},
        ],
        "activities": activities,
        "relationships": relationships,
        "native_p6_verification_required": True,
        "warning": P6_WARNING,
    }


__all__ = [
    "FIXED_QUESTIONS",
    "P6_WARNING",
    "build_schedule_development",
    "build_schedule_development_xer",
    "build_schedule_development_xlsx",
    "project_schedule_to_development_result",
]
