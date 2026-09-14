from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import os
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence, Tuple

from openpyxl import load_workbook


# The editor password's SHA-256 is never committed. Set SCHEDULE_EDITOR_PASSWORD_SHA256, or keep it in
# api/.editor-password.sha256 on the local machine (git-ignored). Without either, password-gated editing is refused.
EDITOR_PASSWORD_FILE = Path(__file__).resolve().parent / ".editor-password.sha256"
MAX_LIBRARY_ROWS = 50_000
MAX_LIBRARY_PAYLOAD_BYTES = 24 * 1024 * 1024
HERE = Path(__file__).resolve().parent
ACTIVITY_CATALOG_PATH = HERE / "activity_catalog.txt"


def is_local_runtime() -> bool:
    return not bool(os.getenv("VERCEL"))


def _editor_password_sha256() -> str:
    value = os.getenv("SCHEDULE_EDITOR_PASSWORD_SHA256", "").strip().lower()
    if not value and EDITOR_PASSWORD_FILE.is_file():
        value = EDITOR_PASSWORD_FILE.read_text(encoding="utf-8").strip().lower()
    return value if re.fullmatch(r"[0-9a-f]{64}", value) else ""


def verify_editor_password(password: str) -> bool:
    expected = _editor_password_sha256()
    if not expected:
        return False
    candidate = hashlib.sha256(password.encode("utf-8")).hexdigest()
    return hmac.compare_digest(candidate, expected)


def _safe_project(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-.")
    return cleaned[:120] or "project"


def _database_path() -> Path:
    base = Path(os.getenv("LOCALAPPDATA") or Path.home() / "AppData" / "Local")
    return base / "ProjectIntelligenceHub" / "schedule_intelligence_libraries.sqlite3"


def _connect() -> sqlite3.Connection:
    if not is_local_runtime():
        raise RuntimeError("Persistent library editing is available only through RUN_LOCAL.bat")
    path = _database_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS library_revisions (
          revision_id INTEGER PRIMARY KEY AUTOINCREMENT,
          project_key TEXT NOT NULL,
          library_kind TEXT NOT NULL,
          revision_number INTEGER NOT NULL,
          changed_at TEXT NOT NULL,
          changed_by TEXT NOT NULL,
          reason TEXT NOT NULL,
          payload_json TEXT NOT NULL,
          payload_sha256 TEXT NOT NULL,
          UNIQUE(project_key, library_kind, revision_number)
        )
        """
    )
    connection.commit()
    return connection


def save_library(kind: str, project_name: str, payload: Dict[str, Any], *, changed_by: str, reason: str) -> Dict[str, Any]:
    if kind not in {"productivity", "activity"}:
        raise ValueError("Unsupported controlled library")
    rows = payload.get("rows")
    if not isinstance(rows, list) or len(rows) > MAX_LIBRARY_ROWS:
        raise ValueError(f"Library rows must be an array containing no more than {MAX_LIBRARY_ROWS:,} records")
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), default=str)
    encoded = serialized.encode("utf-8")
    if len(encoded) > MAX_LIBRARY_PAYLOAD_BYTES:
        raise ValueError("Controlled library payload exceeds the local 24 MB save limit")
    project_key = _safe_project(project_name)
    digest = hashlib.sha256(encoded).hexdigest()
    now = datetime.now(timezone.utc).isoformat()
    with _connect() as connection:
        row = connection.execute(
            "SELECT revision_number, payload_sha256 FROM library_revisions WHERE project_key=? AND library_kind=? ORDER BY revision_number DESC LIMIT 1",
            (project_key, kind),
        ).fetchone()
        if row and row[1] == digest:
            return {"saved": False, "unchanged": True, "revision": int(row[0]), "sha256": digest, "database": str(_database_path())}
        revision = int(row[0]) + 1 if row else 1
        connection.execute(
            "INSERT INTO library_revisions(project_key, library_kind, revision_number, changed_at, changed_by, reason, payload_json, payload_sha256) VALUES(?,?,?,?,?,?,?,?)",
            (project_key, kind, revision, now, changed_by.strip() or "Eng. Ahmed Labib", reason.strip() or "Controlled library update", serialized, digest),
        )
        connection.commit()
    return {"saved": True, "unchanged": False, "revision": revision, "sha256": digest, "database": str(_database_path()), "changed_at": now}


def load_library(kind: str, project_name: str) -> Dict[str, Any]:
    if kind not in {"productivity", "activity"}:
        raise ValueError("Unsupported controlled library")
    if not is_local_runtime() or not _database_path().is_file():
        return {"found": False, "kind": kind, "project_name": project_name}
    with _connect() as connection:
        row = connection.execute(
            "SELECT revision_number, changed_at, changed_by, reason, payload_json, payload_sha256 FROM library_revisions WHERE project_key=? AND library_kind=? ORDER BY revision_number DESC LIMIT 1",
            (_safe_project(project_name), kind),
        ).fetchone()
    if not row:
        return {"found": False, "kind": kind, "project_name": project_name}
    return {
        "found": True,
        "kind": kind,
        "project_name": project_name,
        "revision": int(row[0]),
        "changed_at": row[1],
        "changed_by": row[2],
        "reason": row[3],
        "payload": json.loads(row[4]),
        "sha256": row[5],
    }


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    source = _text(value)
    if source.startswith("="):
        return None
    match = re.search(r"[-+]?\d+(?:[,.]\d+)?", source.replace(",", ""))
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _find(headers: Sequence[str], aliases: Sequence[str]) -> int | None:
    normalized = [re.sub(r"[^a-z0-9]+", " ", item.lower()).strip() for item in headers]
    for alias in aliases:
        needle = re.sub(r"[^a-z0-9]+", " ", alias.lower()).strip()
        for index, header in enumerate(normalized):
            if needle == header or (len(needle) >= 4 and needle in header):
                return index
    return None


def _cell(row: Sequence[Any], index: int | None) -> Any:
    return row[index] if index is not None and index < len(row) else None


def _stable_id(prefix: str, *parts: Any) -> str:
    raw = "|".join(_text(part) for part in parts)
    return f"{prefix}-{hashlib.sha256(raw.encode('utf-8')).hexdigest()[:16].upper()}"


def _formula_tokens(formula: str) -> Dict[str, List[str]]:
    if not formula.startswith("="):
        return {"constants": [], "operators": []}
    return {
        "constants": re.findall(r"(?<![A-Za-z])[-+]?\d+(?:\.\d+)?", formula),
        "operators": re.findall(r"[+\-*/^]", formula[1:]),
    }


def extract_lossless_excel(filename: str, payload: bytes, document_id: str) -> Dict[str, Any]:
    formulas = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    evaluated = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    raw_cells: List[Dict[str, Any]] = []
    raw_rows: List[Dict[str, Any]] = []
    for sheet_index, source_sheet in enumerate(formulas.worksheets, start=1):
        value_sheet = evaluated[source_sheet.title]
        for row_index, source_row in enumerate(source_sheet.iter_rows(), start=1):
            values = [_text(cell.value) for cell in source_row]
            if not any(values):
                continue
            row_id = _stable_id("RAW-XLS-ROW", document_id, source_sheet.title, row_index)
            raw_rows.append({
                "raw_excel_row_id": row_id,
                "document_id": document_id,
                "sheet_name": source_sheet.title,
                "sheet_order": sheet_index,
                "row_number": row_index,
                "original_values": values,
                "source_sequence_number": len(raw_rows) + 1,
            })
            for cell in source_row:
                if cell.value is None or cell.value == "":
                    continue
                formula = _text(cell.value) if cell.data_type == "f" or _text(cell.value).startswith("=") else ""
                evaluated_value = value_sheet[cell.coordinate].value if formula else cell.value
                raw_cells.append({
                    "raw_excel_cell_id": _stable_id("RAW-XLS-CELL", document_id, source_sheet.title, cell.coordinate),
                    "document_id": document_id,
                    "sheet_name": source_sheet.title,
                    "cell_address": cell.coordinate,
                    "row_number": cell.row,
                    "column_number": cell.column,
                    "original_value": cell.value,
                    "original_formula": formula or None,
                    "evaluated_value": evaluated_value,
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "formula_tokens": _formula_tokens(formula),
                })
    return {"raw_excel_rows": raw_rows, "raw_excel_cells": raw_cells}


def _infer_rate_basis(header: str) -> str:
    lowered = header.lower()
    if "daily production" in lowered:
        return "CREW_DAILY_PRODUCTION"
    if "uom / hr" in lowered or "quantity per hour" in lowered:
        return "QUANTITY_PER_CREW_HOUR"
    if "manhour" in lowered or "man-hour" in lowered:
        return "UNKNOWN"
    return "UNKNOWN"


def build_productivity_library(documents: Iterable[Any], packed: Sequence[Tuple[str, bytes]], project_name: str) -> Dict[str, Any]:
    rows: List[Dict[str, Any]] = []
    raw_documents: List[Dict[str, Any]] = []
    raw_excel_rows: List[Dict[str, Any]] = []
    raw_excel_cells: List[Dict[str, Any]] = []
    packed_by_name = {name: payload for name, payload in packed}
    for doc_index, document in enumerate(documents, start=1):
        document_id = _stable_id("DOC", document.filename, document.sha256)
        raw_documents.append({
            "document_id": document_id,
            "original_filename": document.filename,
            "file_type": document.extension,
            "source_project": project_name,
            "file_size": document.size_bytes,
            "hash": document.sha256,
            "parser": document.parser,
            "status": "CATALOGUED" if not document.text and not document.tables else "EXTRACTED",
            "warnings": list(document.warnings),
            "source_sequence_number": doc_index,
        })
        if document.extension in {".xlsx", ".xlsm"} and document.filename in packed_by_name:
            lossless = extract_lossless_excel(document.filename, packed_by_name[document.filename], document_id)
            raw_excel_rows.extend(lossless["raw_excel_rows"])
            raw_excel_cells.extend(lossless["raw_excel_cells"])
        for table_index, table in enumerate(document.tables, start=1):
            headers = [_text(item) for item in table.headers]
            activity_index = _find(headers, ["activity list", "activity description", "description", "work sub-group", "activity", "item"])
            code_index = _find(headers, ["activity code", "code"])
            uom_index = _find(headers, ["uom", "unit"])
            crew_index = _find(headers, ["crew type", "crew"])
            daily_index = _find(headers, ["daily production", "daily output"])
            rate_index = _find(headers, ["actual final rates", "target rate", "productivity rate", "rate"])
            if activity_index is None:
                continue
            for row_index, source_row in enumerate(table.rows, start=2):
                original_name = _text(_cell(source_row, activity_index))
                if not original_name:
                    continue
                daily_source = _cell(source_row, daily_index)
                source_formula = _text(daily_source) if _text(daily_source).startswith("=") else None
                formula_value = None
                if source_formula and daily_index is not None:
                    matched_cell = next((cell for cell in raw_excel_cells if cell.get("document_id") == document_id and cell.get("sheet_name") == table.name and cell.get("row_number") == row_index and cell.get("column_number") == daily_index + 1), None)
                    formula_value = _number(matched_cell.get("evaluated_value")) if matched_cell else None
                daily = formula_value if source_formula else _number(daily_source)
                rate = _number(_cell(source_row, rate_index))
                rate_header = headers[rate_index] if rate_index is not None else headers[daily_index] if daily_index is not None else ""
                record_id = _stable_id("PROD", document_id, table.name, row_index, original_name, _cell(source_row, uom_index), rate, daily)
                rows.append({
                    "id": record_id,
                    "source_record_id": _stable_id("SOURCE", document_id, table.name, row_index),
                    "provenance": "EXCEL_SOURCE" if document.extension in {".xlsx", ".xlsm", ".csv", ".tsv"} else "PDF_SOURCE",
                    "source_document": document.filename,
                    "source_location": f"{table.name} row {row_index}",
                    "source_sequence_number": len(rows) + 1,
                    "original_activity_code": _text(_cell(source_row, code_index)) or None,
                    "original_activity_name": original_name,
                    "display_name": original_name,
                    "original_uom": _text(_cell(source_row, uom_index)) or None,
                    "normalized_uom": None,
                    "original_rate": rate,
                    "original_rate_header": rate_header or None,
                    "original_rate_basis_text": rate_header or None,
                    "normalized_rate_basis": _infer_rate_basis(rate_header),
                    "source_daily_production": daily,
                    "source_formula": source_formula,
                    "source_formula_evaluated_value": formula_value,
                    "crew_type": _text(_cell(source_row, crew_index)) or None,
                    "selected_productivity": None,
                    "quantity": None,
                    "number_of_crews": 1,
                    "calculated_duration": None,
                    "approval_status": "RAW",
                    "working_status": "ACTIVE",
                    "notes": "",
                    "mapping_confidence": 100,
                })
        if document.extension == ".pdf" and document.text:
            for source_line, raw_line in enumerate(document.text.splitlines(), start=1):
                line = raw_line.strip()
                if len(line) < 8 or line.startswith("[PAGE "):
                    continue
                numeric = re.search(r"(?<!\w)([-+]?\d+(?:\.\d+)?)\s*$", line)
                if not numeric:
                    continue
                prefix = line[: numeric.start()].strip(" |\t-")
                if len(prefix) < 5:
                    continue
                rows.append({
                    "id": _stable_id("PROD", document_id, source_line, line),
                    "source_record_id": _stable_id("RAW-PDF", document_id, source_line),
                    "provenance": "PDF_SOURCE",
                    "source_document": document.filename,
                    "source_location": f"extracted text line {source_line}",
                    "source_sequence_number": len(rows) + 1,
                    "original_activity_code": None,
                    "original_activity_name": prefix,
                    "display_name": prefix,
                    "original_uom": None,
                    "normalized_uom": None,
                    "original_rate": float(numeric.group(1)),
                    "original_rate_header": None,
                    "original_rate_basis_text": None,
                    "normalized_rate_basis": "UNKNOWN",
                    "source_daily_production": None,
                    "crew_type": None,
                    "selected_productivity": None,
                    "quantity": None,
                    "number_of_crews": 1,
                    "calculated_duration": None,
                    "approval_status": "REVIEW_REQUIRED",
                    "working_status": "ACTIVE",
                    "notes": "PDF line candidate; verify the source page and rate basis before selection.",
                    "mapping_confidence": 30,
                })
    formula_cells = sum(1 for item in raw_excel_cells if item.get("original_formula"))
    return {
        "kind": "productivity",
        "project_name": project_name,
        "rows": rows[:MAX_LIBRARY_ROWS],
        "raw_documents": raw_documents,
        "raw_excel_rows": raw_excel_rows,
        "raw_excel_cells": raw_excel_cells,
        "import_audit": {
            "source_files": len(raw_documents),
            "source_records_identified": len(rows),
            "source_records_preserved": len(rows),
            "excel_rows_preserved": len(raw_excel_rows),
            "excel_cells_preserved": len(raw_excel_cells),
            "formula_cells_preserved": formula_cells,
            "zero_rates": sum(1 for item in rows if item.get("original_rate") == 0),
            "blank_rates": sum(1 for item in rows if item.get("original_rate") is None and item.get("source_daily_production") is None),
            "unknown_rate_basis": sum(1 for item in rows if item.get("normalized_rate_basis") == "UNKNOWN"),
        },
        "preservation_rule": "Source evidence is immutable. Edits apply to engineered fields and append a revision.",
    }


def load_activity_catalog() -> Dict[str, Any]:
    if not ACTIVITY_CATALOG_PATH.is_file():
        return {"kind": "activity", "rows": [], "error": "Bundled activity catalogue is unavailable"}
    rows: List[Dict[str, Any]] = []
    category = "BUILDING PROJECT"
    for raw in ACTIVITY_CATALOG_PATH.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("===") and line.endswith("==="):
            category = line.strip("= ")
            continue
        sequence = len(rows) + 1
        rows.append({
            "id": _stable_id("ACT", category, sequence, line),
            "source_sequence_number": sequence,
            "source_category": category,
            "source_activity_name": line,
            "display_name": line,
            "project_family": category.replace("ADDITIONAL ", "").replace(" ACTIVITY LIST", "").replace(" ACTIVITIES", "").strip(),
            "discipline": None,
            "subdiscipline": None,
            "work_package": None,
            "system": None,
            "uom": None,
            "approval_status": "RAW",
            "working_status": "ACTIVE",
            "notes": "",
            "provenance": "SUPPLIED_ACTIVITY_CATALOG",
        })
    return {
        "kind": "activity",
        "project_name": "Company activity catalogue",
        "rows": rows,
        "source_snapshot": list(rows),
        "import_audit": {"source_records_identified": len(rows), "source_records_preserved": len(rows)},
        "preservation_rule": "Original source names, spelling, order and categories remain immutable; edits create an engineered display layer.",
    }
